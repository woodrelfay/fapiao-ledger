"""命令行：扫描文件夹 → 预览/执行改名与分类 → 导出 Excel 台账。

默认 dry-run 只预览不改文件；加 --apply 才真正改名，加 --organize 按月份分文件夹。
"""
from __future__ import annotations

import argparse
import csv
import re
import shutil
import sys
from pathlib import Path

from .parser import Invoice, scan_folder

DEFAULT_PATTERN = "{date}_{seller}_{amount}_{invoice_no}"


def unique_total(invoices: list[Invoice]) -> float:
    """价税合计（不含重复）：同一发票号只计一次。

    查重标记会标出同号的全部副本，若简单排除所有被标记的发票，
    一组重复发票的金额会被整组漏掉，所以合计要按发票号去重计一次。
    无发票号的记录无法查重，照常计入。
    """
    seen: set[str] = set()
    total = 0.0
    for inv in invoices:
        if inv.invoice_no:
            if inv.invoice_no in seen:
                continue
            seen.add(inv.invoice_no)
        total += inv.amount_total or 0
    return round(total, 2)


def safe_name(text: str, limit: int = 24) -> str:
    text = re.sub(r'[\\/:*?"<>|\s]+', "", text or "")
    return text[:limit] or "未知"


def build_new_name(inv: Invoice, pattern: str, suffix: str) -> str:
    amount = f"{inv.amount_total:.2f}" if inv.amount_total is not None else "金额未知"
    name = pattern.format(
        date=inv.date or "日期未知",
        seller=safe_name(inv.seller, 20),
        buyer=safe_name(inv.buyer, 20),
        amount=amount,
        invoice_no=inv.invoice_no or "号码未知",
        month=inv.month,
    )
    return safe_name(name, 120) + suffix.lower()


def plan_targets(folder: Path, invoices: list[Invoice], pattern: str,
                 organize: bool) -> list[tuple[Path, Path]]:
    """返回 (原路径, 目标路径) 列表，自动避让重名。"""
    used: set[Path] = set()
    plan: list[tuple[Path, Path]] = []
    sources = sorted(p for p in folder.rglob("*") if p.is_file())
    by_name: dict[str, Path] = {}
    for p in sources:
        by_name.setdefault(p.name, p)
    for inv in invoices:
        src = by_name.get(inv.source_file)
        if src is None:
            continue
        dest_dir = (src.parent / inv.month) if organize else src.parent
        dest = dest_dir / build_new_name(inv, pattern, src.suffix)
        base, n = dest, 2
        while dest in used or (dest.exists() and dest != src):
            dest = base.with_name(f"{base.stem}_{n}{base.suffix}")
            n += 1
        used.add(dest)
        plan.append((src, dest))
    return plan


def export_ledger(invoices: list[Invoice], new_names: dict[str, str], out: Path) -> None:
    rows = []
    for i, inv in enumerate(sorted(invoices, key=lambda x: (x.date, x.invoice_no)), 1):
        rows.append([i, inv.date, inv.invoice_no, inv.seller, inv.buyer,
                     inv.amount_excl, inv.tax, inv.amount_total,
                     inv.source_file, new_names.get(inv.source_file, ""),
                     "是" if inv.is_duplicate else "", inv.note])
    header = ["序号", "开票日期", "发票号码", "销售方", "购买方", "金额(不含税)",
              "税额", "价税合计", "原文件名", "新文件名", "重复", "备注"]
    if out.suffix.lower() == ".csv":
        with out.open("w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(header)
            w.writerows(rows)
        return
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "发票台账"
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    dup_fill = PatternFill("solid", fgColor="FFD7D7")
    for r in rows:
        ws.append(r)
        if r[10] == "是":
            for c in ws[ws.max_row]:
                c.fill = dup_fill
    total = unique_total(invoices)
    ws.append(["", "", "", "", "合计（不含重复）", "", "", total, "", "", "", ""])
    for col, width in zip("ABCDEFGHIJKL", [6, 12, 24, 30, 30, 13, 10, 13, 28, 40, 8, 30]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    wb.save(out)


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(
        prog="fapiao-ledger",
        description="电子发票批量整理：重命名、查重、按月分类、导出 Excel 台账（纯本地运行）")
    ap.add_argument("folder", help="存放发票的文件夹（支持 PDF / XML / OFD，含子文件夹）")
    ap.add_argument("--apply", action="store_true", help="真正执行改名（默认只预览）")
    ap.add_argument("--organize", action="store_true", help="按月份分子文件夹（自动启用 --apply）")
    ap.add_argument("--pattern", default=DEFAULT_PATTERN,
                    help=f"改名模板，可用 {{date}} {{seller}} {{buyer}} {{amount}} {{invoice_no}} {{month}}，默认 {DEFAULT_PATTERN}")
    ap.add_argument("--export", default="", help="台账导出路径（.xlsx 或 .csv），默认 <文件夹>/发票台账.xlsx")
    args = ap.parse_args(argv)

    folder = Path(args.folder).expanduser()
    if not folder.is_dir():
        print(f"文件夹不存在：{folder}")
        return 1

    invoices = scan_folder(folder)
    if not invoices:
        print("没有找到 PDF / XML / OFD 发票文件。")
        return 0

    plan = plan_targets(folder, invoices, args.pattern, args.organize)
    new_names = {src.name: dst.name for src, dst in plan}

    print(f"共扫描 {len(invoices)} 张发票")
    for src, dst in plan:
        print(f"  {src.name}  →  {dst.relative_to(folder) if args.organize else dst.name}")
    dups = [i for i in invoices if i.is_duplicate]
    failed = [i for i in invoices if not i.parse_ok]
    total = unique_total(invoices)
    print(f"解析成功 {len(invoices) - len(failed)} 张｜重复 {len(dups)} 张｜"
          f"价税合计（不含重复）¥{total:,.2f}")
    for inv in failed:
        print(f"  ⚠ {inv.source_file}：{inv.note}")

    do_apply = args.apply or args.organize
    if do_apply:
        log_rows = []
        for src, dst in plan:
            if src != dst:
                dst.parent.mkdir(parents=True, exist_ok=True)
                shutil.move(str(src), str(dst))
                log_rows.append([str(src), str(dst)])
        if log_rows:
            log_path = folder / "rename_log.csv"
            with log_path.open("w", newline="", encoding="utf-8-sig") as f:
                w = csv.writer(f)
                w.writerow(["原路径", "新路径"])
                w.writerows(log_rows)
            print(f"已改名 {len(log_rows)} 个文件，记录在 {log_path.name}（可据此还原）")
    else:
        print("当前是预览模式，没有改动任何文件。确认无误后加 --apply 执行。")

    out = Path(args.export) if args.export else folder / "发票台账.xlsx"
    export_ledger(invoices, new_names, out)
    print(f"台账已导出：{out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
